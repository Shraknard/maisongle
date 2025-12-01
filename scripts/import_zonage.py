"""Script to import zonage ABC data from Excel file into database."""
import sys
import os

# Add parent directory to path for imports
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook
from app.database import SessionLocal, engine, Base
from app.models.zonage import ZonageCommune


def import_zonage_data(excel_path: str):
    """Import zonage ABC data from Excel file.
    
    Expected columns (based on actual file):
    - CODGEO: INSEE code
    - DEP: Department code
    - LIBGEO: Commune name
    - Zonage en vigueur depuis le 5 septembre 2025: Zone (A, Abis, B1, B2, C)
    """
    
    # Create tables if they don't exist
    Base.metadata.create_all(bind=engine)
    
    print(f"Loading Excel file: {excel_path}")
    wb = load_workbook(excel_path, read_only=True)
    ws = wb.active
    
    db = SessionLocal()
    
    try:
        # Get header row
        headers = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(h).strip() if h else "" for h in row]
            break
        
        print(f"Headers found: {headers}")
        
        # Map columns based on actual file structure
        # CODGEO=0, DEP=1, LIBGEO=2, Zonage=3
        col_insee = 0  # CODGEO
        col_dept = 1   # DEP
        col_nom = 2    # LIBGEO
        col_zonage = 3 # Zonage en vigueur
        
        # Clear existing data
        db.query(ZonageCommune).delete()
        db.commit()
        
        # Import data
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            insee = row[col_insee] if len(row) > col_insee else None
            zonage = row[col_zonage] if len(row) > col_zonage else None
            
            if not insee or not zonage:
                continue
            
            commune = ZonageCommune(
                code_insee=str(insee).strip(),
                nom_commune=str(row[col_nom]).strip() if len(row) > col_nom and row[col_nom] else None,
                code_departement=str(row[col_dept]).strip() if len(row) > col_dept and row[col_dept] else None,
                zonage=str(zonage).strip(),
            )
            db.add(commune)
            count += 1
            
            # Commit in batches
            if count % 1000 == 0:
                db.commit()
                print(f"Imported {count} communes...")
        
        db.commit()
        print(f"Successfully imported {count} communes")
        
    except Exception as e:
        print(f"Error importing data: {e}")
        db.rollback()
        raise
    finally:
        db.close()
        wb.close()


if __name__ == "__main__":
    # Default path to Excel file
    excel_path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "data",
        "zonage_abc.xlsx"
    )
    
    # Allow override via command line
    if len(sys.argv) > 1:
        excel_path = sys.argv[1]
    
    import_zonage_data(excel_path)
